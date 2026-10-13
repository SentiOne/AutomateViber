from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from .io import CliError, iter_message_records
from .models import MessageRecord


@dataclass(frozen=True)
class BlockRun:
    internal_session_id: str
    block_id: str
    rows: tuple[MessageRecord, ...]
    previous_row_bot_messages: tuple[tuple[str, ...], ...]
    previous_block_id: str | None
    next_block_id: str | None
    previous_run_last_message: str


@dataclass(frozen=True)
class BlockDrilldownData:
    target_block_id: str
    target_block_name: str
    target_context_name: str
    total_sessions: int
    sessions_with_target_block: int
    total_target_rows: int
    total_target_runs: int
    runs_with_previous_block: int
    runs_with_next_block: int
    runs_with_fallback: int
    fallback_turns: int
    previous_user_messages: Counter[str]
    phrases_to_current: Counter[str]
    phrases_to_next: Counter[str]
    fallback_phrases: Counter[str]
    previous_bot_messages: Counter[str]
    next_bot_messages: Counter[str]
    fallback_bot_messages: Counter[str]
    next_blocks: Counter[str]
    previous_blocks: Counter[str]
    block_display_names: dict[str, str]
    block_context_names: dict[str, str]


def _clean_phrase(value: str) -> str:
    return value.strip()


def _preferred_name(name_counts: Counter[str]) -> str:
    return sorted(name_counts.items(), key=lambda item: (-item[1], item[0]))[0][0]


def _preferred_context_name(context_counts: Counter[str]) -> str:
    if not context_counts:
        return ""
    return sorted(context_counts.items(), key=lambda item: (-item[1], item[0]))[0][0]


def build_block_drilldown_data(zip_path: Path, target_block_id: str) -> BlockDrilldownData:
    session_rows: defaultdict[str, list[MessageRecord]] = defaultdict(list)
    block_name_counts: defaultdict[str, Counter[str]] = defaultdict(Counter)
    block_context_counts: defaultdict[str, Counter[str]] = defaultdict(Counter)

    for record in iter_message_records(zip_path):
        session_rows[record.internal_session_id].append(record)
        if record.dialog_state is not None:
            block_name_counts[record.dialog_state.block_id][record.dialog_state.block_name] += 1
            if record.dialog_state.context:
                block_context_counts[record.dialog_state.block_id][record.dialog_state.context] += 1

    if target_block_id not in block_name_counts:
        raise CliError(f"Block id not found in messages.json: {target_block_id}")

    block_display_names = {
        block_id: _preferred_name(name_counts)
        for block_id, name_counts in block_name_counts.items()
    }
    block_context_names = {
        block_id: _preferred_context_name(block_context_counts[block_id])
        for block_id in block_name_counts
    }

    target_runs: list[BlockRun] = []

    for internal_session_id, rows in session_rows.items():
        rows.sort(key=lambda row: (row.time, row.row_id))
        previous_bot_messages_by_row_id: dict[int, tuple[str, ...]] = {}
        previous_row: MessageRecord | None = None
        for row in rows:
            previous_bot_messages_by_row_id[row.row_id] = previous_row.bot_messages if previous_row is not None else ()
            previous_row = row

        non_null_rows = [row for row in rows if row.dialog_state is not None]
        if not non_null_rows:
            continue

        grouped_runs: list[tuple[str, list[MessageRecord]]] = []
        current_block_id: str | None = None
        current_rows: list[MessageRecord] = []

        for row in non_null_rows:
            block_id = row.dialog_state.block_id
            if block_id == current_block_id:
                current_rows.append(row)
                continue

            if current_block_id is not None:
                grouped_runs.append((current_block_id, current_rows))

            current_block_id = block_id
            current_rows = [row]

        if current_block_id is not None:
            grouped_runs.append((current_block_id, current_rows))

        for index, (block_id, run_rows) in enumerate(grouped_runs):
            if block_id != target_block_id:
                continue

            previous_block_id = grouped_runs[index - 1][0] if index > 0 else None
            next_block_id = grouped_runs[index + 1][0] if index + 1 < len(grouped_runs) else None
            target_runs.append(
                BlockRun(
                    internal_session_id=internal_session_id,
                    block_id=block_id,
                    rows=tuple(run_rows),
                    previous_row_bot_messages=tuple(
                        previous_bot_messages_by_row_id[row.row_id]
                        for row in run_rows
                    ),
                    previous_block_id=previous_block_id,
                    next_block_id=next_block_id,
                    previous_run_last_message=grouped_runs[index - 1][1][-1].message if index > 0 else "",
                )
            )

    previous_user_messages: Counter[str] = Counter()
    phrases_to_current: Counter[str] = Counter()
    phrases_to_next: Counter[str] = Counter()
    fallback_phrases: Counter[str] = Counter()
    previous_bot_messages: Counter[str] = Counter()
    next_bot_messages: Counter[str] = Counter()
    fallback_bot_messages: Counter[str] = Counter()
    next_blocks: Counter[str] = Counter()
    previous_blocks: Counter[str] = Counter()

    sessions_with_target_block = len({run.internal_session_id for run in target_runs})
    total_target_rows = 0
    runs_with_previous_block = 0
    runs_with_next_block = 0
    runs_with_fallback = 0
    fallback_turns = 0

    for run in target_runs:
        total_target_rows += len(run.rows)

        for row in run.rows:
            for bot_message in row.bot_messages:
                phrase = _clean_phrase(bot_message)
                if phrase:
                    next_bot_messages[phrase] += 1

        for row in run.rows:
            phrase = _clean_phrase(row.message)
            if phrase:
                phrases_to_current[phrase] += 1

        if run.previous_block_id is not None:
            runs_with_previous_block += 1
            previous_blocks[run.previous_block_id] += 1
            previous_message = _clean_phrase(run.previous_run_last_message)
            if previous_message:
                previous_user_messages[previous_message] += 1

        for previous_row_bot_messages in run.previous_row_bot_messages:
            for bot_message in previous_row_bot_messages:
                cleaned_bot_message = _clean_phrase(bot_message)
                if cleaned_bot_message:
                    previous_bot_messages[cleaned_bot_message] += 1

        if run.next_block_id is not None:
            runs_with_next_block += 1
            next_blocks[run.next_block_id] += 1

            last_phrase = _clean_phrase(run.rows[-1].message)
            if last_phrase:
                phrases_to_next[last_phrase] += 1

        if len(run.rows) > 1:
            runs_with_fallback += 1
            fallback_turns += len(run.rows) - 1
            for row in run.rows[:-1]:
                phrase = _clean_phrase(row.message)
                if phrase:
                    fallback_phrases[phrase] += 1
                for bot_message in row.bot_messages:
                    cleaned_bot_message = _clean_phrase(bot_message)
                    if cleaned_bot_message:
                        fallback_bot_messages[cleaned_bot_message] += 1

    return BlockDrilldownData(
        target_block_id=target_block_id,
        target_block_name=block_display_names[target_block_id],
        target_context_name=block_context_names[target_block_id],
        total_sessions=len(session_rows),
        sessions_with_target_block=sessions_with_target_block,
        total_target_rows=total_target_rows,
        total_target_runs=len(target_runs),
        runs_with_previous_block=runs_with_previous_block,
        runs_with_next_block=runs_with_next_block,
        runs_with_fallback=runs_with_fallback,
        fallback_turns=fallback_turns,
        previous_user_messages=previous_user_messages,
        phrases_to_current=phrases_to_current,
        phrases_to_next=phrases_to_next,
        fallback_phrases=fallback_phrases,
        previous_bot_messages=previous_bot_messages,
        next_bot_messages=next_bot_messages,
        fallback_bot_messages=fallback_bot_messages,
        next_blocks=next_blocks,
        previous_blocks=previous_blocks,
        block_display_names=block_display_names,
        block_context_names=block_context_names,
    )


def _style_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.auto_filter.ref = sheet.dimensions


def _write_counter_sheet(sheet, header_name: str, values: Counter[str]) -> None:
    sheet.append([header_name, "Occurrences"])
    for phrase, occurrences in sorted(values.items(), key=lambda item: (-item[1], item[0])):
        sheet.append([phrase, occurrences])
    _style_sheet(sheet)
    sheet.column_dimensions["A"].width = 80
    sheet.column_dimensions["B"].width = 14


def _write_block_stats_sheet(
    sheet,
    prefix: str,
    values: Counter[str],
    display_names: dict[str, str],
    context_names: dict[str, str],
) -> None:
    total = sum(values.values())
    sheet.append([f"{prefix} block id", f"{prefix} block name", f"{prefix} context name", "Occurrences", "Percentage"])
    for block_id, occurrences in sorted(values.items(), key=lambda item: (-item[1], item[0])):
        percentage = (occurrences / total) if total else 0.0
        sheet.append([block_id, display_names.get(block_id, ""), context_names.get(block_id, ""), occurrences, percentage])
    _style_sheet(sheet)
    sheet.column_dimensions["A"].width = 40
    sheet.column_dimensions["B"].width = 40
    sheet.column_dimensions["C"].width = 32
    sheet.column_dimensions["D"].width = 14
    sheet.column_dimensions["E"].width = 14
    for cell in sheet["E"][1:]:
        cell.number_format = "0.00%"


def _write_overview_sheet(sheet, data: BlockDrilldownData) -> None:
    rows = [
        ("Target block id", data.target_block_id, "Exact dialogState.id selected for this drilldown report."),
        ("Target block name", data.target_block_name, "Most frequent observed display name for the selected block id."),
        ("Target context name", data.target_context_name, "Most frequent observed context name for the selected block id."),
        ("Total sessions", data.total_sessions, "Number of sessions found in messages.json."),
        ("Sessions with target block", data.sessions_with_target_block, "Unique sessions where the selected block id appeared at least once."),
        ("Total target turns", data.total_target_rows, "Raw turns from messages.json whose dialogState.id equals the selected block id."),
        ("Total target runs", data.total_target_runs, "Consecutive groups of user turns with the selected dialogState.id after collapsing repeated adjacent turns."),
        ("Runs with previous block", data.runs_with_previous_block, "Target runs that have a different non-null block immediately before them."),
        ("Runs with next block", data.runs_with_next_block, "Target runs that have a different non-null block immediately after them."),
        ("Runs with fallback", data.runs_with_fallback, "Target runs whose length is greater than one, meaning the same input-requesting dialogState.id repeated across consecutive user turns. This is a fallback heuristic, not proof of exact runtime state."),
        ("Fallback turns", data.fallback_turns, "Total extra repeated dialogState.id turns inside all target runs, calculated as sum(run_length - 1)."),
        ("Unique previous user messages", len(data.previous_user_messages), "Different user messages that appeared just before this block became active."),
        ("Unique bot messages", len(data.previous_bot_messages), "Different bot messages that users replied to while they were in this block."),
        ("Unique user replies", len(data.phrases_to_current), "Different user messages sent as replies to this block."),
        ("Unique next bot messages", len(data.next_bot_messages), "Different bot messages sent after those user replies."),
        ("Unique non-fallback user messages", len(data.phrases_to_next), "Different user replies that were followed by a different next block."),
        ("Unique fallback user messages", len(data.fallback_phrases), "Different user replies after which the next exported user turn still referenced the same dialogState.id."),
        ("Unique fallback bot replies", len(data.fallback_bot_messages), "Different bot replies emitted after user replies in repeated same-dialogState runs."),
    ]
    glossary_rows = [
        ("Target", "If target = A, analyze only turns and runs for block A.", "The selected block id being analyzed in this report."),
        ("Turn", "A -> A -> B contains 3 raw turns.", "One raw record from messages.json where dialogState.id may equal the target block id."),
        ("Run", "A -> A -> B contains 2 runs: [A -> A] and [B].", "One consecutive group of adjacent user turns in the same session that share the same dialogState.id."),
    ]
    sheet.append(["Metric", "Value", "Description"])
    for metric, value, description in rows:
        sheet.append([metric, value, description])
    sheet.append(["", "", ""])
    sheet.append(["Glossary term", "Example", "Description"])
    for term, example, description in glossary_rows:
        sheet.append([term, example, description])
    _style_sheet(sheet)
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 28
    sheet.column_dimensions["C"].width = 96


def write_block_overview_report(data: BlockDrilldownData, output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Overview"
    _write_overview_sheet(sheet, data)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_block_phrases_to_next_report(data: BlockDrilldownData, output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Non-fallback user messages"
    _write_counter_sheet(sheet, "User message", data.phrases_to_next)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_block_fallback_phrases_report(data: BlockDrilldownData, output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fallback user messages"
    _write_counter_sheet(sheet, "User message", data.fallback_phrases)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_block_next_blocks_report(data: BlockDrilldownData, output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Next Blocks"
    _write_block_stats_sheet(sheet, "Next", data.next_blocks, data.block_display_names, data.block_context_names)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_block_previous_blocks_report(data: BlockDrilldownData, output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Previous Blocks"
    _write_block_stats_sheet(sheet, "Previous", data.previous_blocks, data.block_display_names, data.block_context_names)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_block_phrases_to_current_report(data: BlockDrilldownData, output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "User replies"
    _write_counter_sheet(sheet, "User reply", data.phrases_to_current)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_block_bot_messages_report(data: BlockDrilldownData, output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Next bot messages"
    _write_counter_sheet(sheet, "Bot message", data.next_bot_messages)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_block_previous_bot_messages_report(data: BlockDrilldownData, output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bot messages"
    _write_counter_sheet(sheet, "Bot message", data.previous_bot_messages)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_block_previous_user_messages_report(data: BlockDrilldownData, output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Previous user messages"
    _write_counter_sheet(sheet, "User message", data.previous_user_messages)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_block_drilldown_report(data: BlockDrilldownData, output_path: Path) -> None:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    _write_overview_sheet(overview, data)

    sheet = workbook.create_sheet("Previous Blocks")
    _write_block_stats_sheet(sheet, "Previous", data.previous_blocks, data.block_display_names, data.block_context_names)

    sheet = workbook.create_sheet("Previous user messages")
    _write_counter_sheet(sheet, "User message", data.previous_user_messages)

    sheet = workbook.create_sheet("Bot messages")
    _write_counter_sheet(sheet, "Bot message", data.previous_bot_messages)

    sheet = workbook.create_sheet("User replies")
    _write_counter_sheet(sheet, "User reply", data.phrases_to_current)

    sheet = workbook.create_sheet("Next bot messages")
    _write_counter_sheet(sheet, "Bot message", data.next_bot_messages)

    sheet = workbook.create_sheet("Non-fallback user messages")
    _write_counter_sheet(sheet, "User message", data.phrases_to_next)

    sheet = workbook.create_sheet("Fallback user messages")
    _write_counter_sheet(sheet, "User message", data.fallback_phrases)

    sheet = workbook.create_sheet("Fallback bot replies")
    _write_counter_sheet(sheet, "Bot reply", data.fallback_bot_messages)

    sheet = workbook.create_sheet("Next Blocks")
    _write_block_stats_sheet(sheet, "Next", data.next_blocks, data.block_display_names, data.block_context_names)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
