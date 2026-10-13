from __future__ import annotations

import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from scripts.debugger_lib.fallback_circular_report import create_fallback_circular_report


class FallbackCircularReportTests(unittest.TestCase):
    def test_report_buckets_same_block_runs_into_1_2_and_3_plus(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            base = Path(tmp_dir)
            zip_path = base / "conversations.zip"
            output_path = base / "fallback-circular.xlsx"

            message_rows = [
                {"internalSessionId": "sid-1", "time": 100, "dialogState": {"id": "a", "name": "say:A", "context": "Main"}},
                {"internalSessionId": "sid-1", "time": 200, "dialogState": {"id": "a", "name": "say:A", "context": "Main"}},
                {"internalSessionId": "sid-1", "time": 300, "dialogState": {"id": "b", "name": "say:B", "context": "Exit"}},
                {"internalSessionId": "sid-2", "time": 100, "dialogState": {"id": "a", "name": "say:A", "context": "Main"}},
                {"internalSessionId": "sid-2", "time": 200, "dialogState": {"id": "a", "name": "say:A", "context": "Main"}},
                {"internalSessionId": "sid-2", "time": 300, "dialogState": {"id": "a", "name": "say:A", "context": "Main"}},
                {"internalSessionId": "sid-3", "time": 100, "dialogState": {"id": "a", "name": "say:A", "context": "Main"}},
                {"internalSessionId": "sid-3", "time": 200, "dialogState": {"id": "a", "name": "say:A", "context": "Main"}},
                {"internalSessionId": "sid-3", "time": 300, "dialogState": {"id": "a", "name": "say:A", "context": "Main"}},
                {"internalSessionId": "sid-3", "time": 400, "dialogState": {"id": "a", "name": "say:A", "context": "Main"}},
                {"internalSessionId": "sid-4", "time": 100, "dialogState": {"id": "b", "name": "say:B", "context": "Exit"}},
                {"internalSessionId": "sid-5", "time": 100, "dialogState": None},
            ]

            with zipfile.ZipFile(zip_path, "w") as archive:
                archive.writestr(
                    "messages.json",
                    "".join(json.dumps(row) + "\n" for row in message_rows),
                )

            summary = create_fallback_circular_report(zip_path, output_path)

            self.assertEqual(summary.total_sessions, 5)
            self.assertEqual(summary.sessions_with_block_visits, 4)
            self.assertEqual(summary.sessions_with_fallback_episodes, 3)

            workbook = load_workbook(output_path)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))

            self.assertEqual(
                rows,
                [
                    (
                        "Block name",
                        "Block id",
                        "Context name",
                        "Visits",
                        "Fallbacks: 1",
                        "Fallbacks: 2",
                        "Fallbacks: 3+",
                        "Total fallback episodes",
                    ),
                    ("say:A", "a", "Main", 3, 1, 1, 1, 3),
                    ("say:B", "b", "Exit", 2, 0, 0, 0, 0),
                ],
            )


if __name__ == "__main__":
    unittest.main()
