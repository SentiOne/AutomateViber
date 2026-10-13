from __future__ import annotations

import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from scripts.debugger_lib.tag_conversation_graph_report import (
    build_tag_conversation_graph_data,
    create_tag_conversation_graph_html_report,
    create_tag_conversation_graph_report,
    create_tag_drop_terminal_report,
)


class TagConversationGraphReportTests(unittest.TestCase):
    def _make_zip(self, zip_path: Path) -> None:
        message_rows = [
            {"id": 1, "internalSessionId": "sid-1", "time": 100, "message": "hello", "tags": ["entry", "faq"], "dialogState": None},
            {"id": 2, "internalSessionId": "sid-1", "time": 200, "message": "next", "tags": ["faq", "handoff"], "commands": [{"type": "Disconnect", "params": None}], "dialogState": None},
            {"id": 3, "internalSessionId": "sid-2", "time": 100, "message": "start", "tags": ["entry", "faq"], "dialogState": None},
            {"id": 4, "internalSessionId": "sid-2", "time": 200, "message": "", "tags": ["done"], "dialogState": None},
            {"id": 5, "internalSessionId": "sid-3", "time": 100, "message": "only", "tags": ["entry"], "commands": [{"type": "redirect", "params": None}], "dialogState": None},
            {"id": 6, "internalSessionId": "sid-4", "time": 100, "message": "wait", "tags": ["wait"], "dialogState": None},
            {"id": 7, "internalSessionId": "sid-4", "time": 200, "message": "follow up", "tags": [], "dialogState": None},
            {"id": 8, "internalSessionId": "sid-5", "time": 100, "message": "", "tags": [], "dialogState": None},
        ]

        with zipfile.ZipFile(zip_path, "w") as archive:
            archive.writestr(
                "messages.json",
                "".join(json.dumps(row) + "\n" for row in message_rows),
            )

    def test_build_tag_conversation_graph_data_uses_ordered_tag_sequence(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            zip_path = Path(tmp_dir) / "conversations.zip"
            self._make_zip(zip_path)

            data = build_tag_conversation_graph_data(zip_path)

            self.assertEqual(data.total_sessions, 5)
            self.assertEqual(data.total_nodes, 6)
            self.assertEqual(data.total_edges, 5)
            self.assertEqual(sorted(data.nodes), ["done", "entry", "faq", "handoff", "wait"])
            self.assertEqual(data.nodes["entry"].occurrences, 3)
            self.assertEqual(data.nodes["entry"].runs, 3)
            self.assertEqual(data.nodes["entry"].sessions, 3)
            self.assertAlmostEqual(data.nodes["entry"].drop_rate, 0.0)
            self.assertAlmostEqual(data.nodes["entry"].terminal_rate, 1 / 3)
            self.assertEqual(data.nodes["faq"].occurrences, 3)
            self.assertEqual(data.nodes["faq"].runs, 2)
            self.assertAlmostEqual(data.nodes["faq"].fallback_rate, 1 / 2)
            self.assertAlmostEqual(data.nodes["handoff"].drop_rate, 0.0)
            self.assertAlmostEqual(data.nodes["handoff"].terminal_rate, 1.0)
            self.assertAlmostEqual(data.nodes["wait"].drop_rate, 0.0)
            self.assertAlmostEqual(data.nodes["wait"].terminal_rate, 0.0)

            edge_map = {
                (edge.source_tag, edge.destination_tag): edge
                for edge in data.edges
            }
            self.assertAlmostEqual(edge_map[("__start__", "entry")].outgoing_percentage, 3 / 5)
            self.assertAlmostEqual(edge_map[("__start__", "wait")].outgoing_percentage, 1 / 5)
            self.assertAlmostEqual(edge_map[("entry", "faq")].outgoing_percentage, 2 / 3)
            self.assertAlmostEqual(edge_map[("faq", "handoff")].outgoing_percentage, 1 / 2)
            self.assertAlmostEqual(edge_map[("faq", "done")].outgoing_percentage, 1 / 2)

    def test_filtered_tag_graph_restores_orphan_incoming_edges(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            zip_path = Path(tmp_dir) / "conversations.zip"
            self._make_zip(zip_path)

            data = build_tag_conversation_graph_data(
                zip_path,
                min_node_occurrences=1,
                min_edge_count=1,
                min_edge_percent=0.0,
                top_outgoing=1,
            )

            edge_map = {
                (edge.source_tag, edge.destination_tag): edge
                for edge in data.edges
            }
            self.assertIn(("faq", "done"), edge_map)
            self.assertNotIn(("__start__", "done"), edge_map)

    def test_create_tag_conversation_graph_report_writes_svg_and_dot(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            base = Path(tmp_dir)
            zip_path = base / "conversations.zip"
            output_path = base / "tag-graph.svg"
            dot_output_path = base / "tag-graph.dot"
            self._make_zip(zip_path)

            summary = create_tag_conversation_graph_report(
                zip_path,
                output_path,
                dot_output_path=dot_output_path,
            )

            self.assertEqual(summary.total_sessions, 5)
            self.assertEqual(summary.total_nodes, 6)
            self.assertEqual(summary.total_edges, 5)
            self.assertEqual(summary.filtered_nodes, 6)
            self.assertEqual(summary.filtered_edges, 5)
            self.assertTrue(output_path.exists())
            self.assertTrue(dot_output_path.exists())
            self.assertIn("<svg", output_path.read_text(encoding="utf-8"))

            dot_content = dot_output_path.read_text(encoding="utf-8")
            self.assertIn('"__start__" [label="Start\\nSessions: 5"', dot_content)
            self.assertIn('label="entry\\nOccurrences: 3 | Runs: 3\\nFallback: 0.0%\\nDrop: 0.0% | Terminal: 33.3%"', dot_content)
            self.assertIn('label="60.0% (3)"', dot_content)
            self.assertIn('Drop: 0.0% | Terminal: 0.0%', dot_content)

    def test_create_tag_conversation_graph_html_report_writes_interactive_html(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            base = Path(tmp_dir)
            zip_path = base / "conversations.zip"
            output_path = base / "tag-graph.html"
            self._make_zip(zip_path)

            summary = create_tag_conversation_graph_html_report(zip_path, output_path)

            self.assertEqual(summary.total_sessions, 5)
            self.assertEqual(summary.total_nodes, 6)
            self.assertEqual(summary.total_edges, 5)
            html = output_path.read_text(encoding="utf-8")
            self.assertIn("Conversation Tag Graph", html)
            self.assertIn("cytoscape@3.34.0", html)
            self.assertIn("Open+Sans", html)
            self.assertIn("#7C6DFD", html)
            self.assertIn('aria-label="SentiOne Automate logo"', html)
            self.assertIn('"title": "faq"', html)
            self.assertIn('"label": "50.0% (1)"', html)
            self.assertIn('"dropRate": 0.0', html)
            self.assertIn('"terminalRate": 33.33333333333333', html)
            self.assertIn('"Drop rate": "0.0%"', html)
            self.assertIn('"Terminal rate": "33.3%"', html)
            self.assertIn('id="node-border-legend"', html)
            self.assertIn("Tag drop rate border", html)
            self.assertIn("borderRateColor(borderRate)", html)
            self.assertIn('no later user message and no Disconnect or Redirect command', html)

    def test_create_tag_drop_terminal_report_writes_xlsx(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            base = Path(tmp_dir)
            zip_path = base / "conversations.zip"
            output_path = base / "tag-drop-terminal.xlsx"
            self._make_zip(zip_path)

            summary = create_tag_drop_terminal_report(zip_path, output_path)

            self.assertEqual(summary.total_sessions, 5)
            self.assertEqual(summary.total_tags, 5)
            self.assertTrue(output_path.exists())

            workbook = load_workbook(output_path)
            self.assertEqual(workbook.sheetnames, ["Tag Drop Terminal"])
            rows = list(workbook.active.iter_rows(values_only=True))

            self.assertEqual(
                rows,
                [
                    (
                        "Tag",
                        "Occurrences",
                        "Runs",
                        "Sessions",
                        "Sessions percentage of all sessions",
                        "Drop runs",
                        "Drop rate",
                        "Terminal sessions",
                        "Terminal sessions percentage of all sessions",
                        "Terminal rate",
                    ),
                    ("done", 1, 1, 1, 1 / 5, 1, 1, 1, 1 / 5, 1),
                    ("handoff", 1, 1, 1, 1 / 5, 0, 0, 1, 1 / 5, 1),
                    ("entry", 3, 3, 3, 3 / 5, 0, 0, 1, 1 / 5, 1 / 3),
                    ("faq", 3, 2, 2, 2 / 5, 0, 0, 0, 0, 0),
                    ("wait", 1, 1, 1, 1 / 5, 0, 0, 0, 0, 0),
                ],
            )


if __name__ == "__main__":
    unittest.main()
