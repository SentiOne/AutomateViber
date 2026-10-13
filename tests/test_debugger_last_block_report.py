from __future__ import annotations

import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from scripts.debugger_lib.last_block_report import create_last_block_report


class LastBlockReportTests(unittest.TestCase):
    def test_report_uses_sessions_dialog_state_and_skips_nulls(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            base = Path(tmp_dir)
            zip_path = base / "conversations.zip"
            output_path = base / "report.xlsx"

            session_rows = [
                {
                    "internalSessionId": "sid-1",
                    "dialogState": {"id": "block-2", "name": "say:Fallback", "context": "Fallback context"},
                },
                {
                    "internalSessionId": "sid-2",
                    "dialogState": {"id": "block-1", "name": "say:Main", "context": "Main context"},
                },
                {
                    "internalSessionId": "sid-3",
                    "dialogState": {"id": "block-1", "name": "say:Main", "context": "Main context"},
                },
                {
                    "internalSessionId": "sid-4",
                    "dialogState": None,
                },
            ]
            message_rows = [
                {
                    "internalSessionId": "sid-4",
                    "time": 123,
                    "dialogState": {"id": "ignored", "name": "say:Ignored"},
                }
            ]

            with zipfile.ZipFile(zip_path, "w") as archive:
                archive.writestr(
                    "sessions.json",
                    "".join(json.dumps(row) + "\n" for row in session_rows),
                )
                archive.writestr(
                    "messages.json",
                    "".join(json.dumps(row) + "\n" for row in message_rows),
                )

            summary = create_last_block_report(zip_path, output_path, percentage_base="all")

            self.assertEqual(summary.total_sessions, 4)
            self.assertEqual(summary.resolved_sessions, 3)
            self.assertEqual(summary.unresolved_sessions, 1)

            workbook = load_workbook(output_path)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))

            self.assertEqual(
                rows,
                [
                    (
                        "Block name",
                        "Block id",
                        "Context name",
                        "Occurrences as last block",
                        "Overall percentage",
                    ),
                    ("say:Main", "block-1", "Main context", 2, 0.5),
                    ("say:Fallback", "block-2", "Fallback context", 1, 0.25),
                ],
            )

    def test_report_can_use_resolved_percentage_base(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            base = Path(tmp_dir)
            zip_path = base / "conversations.zip"
            output_path = base / "report.xlsx"

            session_rows = [
                {"internalSessionId": "sid-1", "dialogState": {"id": "block-1", "name": "say:Main"}},
                {"internalSessionId": "sid-2", "dialogState": {"id": "block-1", "name": "say:Main"}},
                {"internalSessionId": "sid-3", "dialogState": None},
            ]

            with zipfile.ZipFile(zip_path, "w") as archive:
                archive.writestr(
                    "sessions.json",
                    "".join(json.dumps(row) + "\n" for row in session_rows),
                )

            create_last_block_report(zip_path, output_path, percentage_base="resolved")

            workbook = load_workbook(output_path)
            sheet = workbook.active
            data_row = list(sheet.iter_rows(values_only=True))[1]
            self.assertEqual(data_row, ("say:Main", "block-1", None, 2, 1))


if __name__ == "__main__":
    unittest.main()
