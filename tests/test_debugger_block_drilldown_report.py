from __future__ import annotations

import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from scripts.debugger_lib.block_drilldown_report import (
    build_block_drilldown_data,
    write_block_bot_messages_report,
    write_block_drilldown_report,
    write_block_next_blocks_report,
    write_block_previous_bot_messages_report,
    write_block_previous_user_messages_report,
)


class BlockDrilldownReportTests(unittest.TestCase):
    def _make_zip(self, zip_path: Path) -> None:
        def state(block_id: str, name: str, context: str) -> dict[str, str]:
            return {"id": block_id, "name": name, "context": context}

        message_rows = [
            {"id": 1, "internalSessionId": "sid-1", "time": 100, "message": "", "botMessages": ["welcome"], "dialogState": None},
            {"id": 2, "internalSessionId": "sid-1", "time": 200, "message": "hello", "botMessages": ["ask-a"], "dialogState": state("a", "say:A", "Main")},
            {"id": 3, "internalSessionId": "sid-1", "time": 300, "message": "still help", "botMessages": ["clarify-a"], "dialogState": state("a", "say:A renamed", "Main")},
            {"id": 4, "internalSessionId": "sid-1", "time": 400, "message": "bye", "botMessages": ["to-b"], "dialogState": state("b", "say:B", "Exit")},
            {"id": 5, "internalSessionId": "sid-2", "time": 100, "message": "toA", "botMessages": ["from-c"], "dialogState": state("c", "say:C", "Router")},
            {"id": 6, "internalSessionId": "sid-2", "time": 200, "message": "question", "botMessages": ["ask-a"], "dialogState": state("a", "say:A", "Main")},
            {"id": 7, "internalSessionId": "sid-2", "time": 300, "message": "goD", "botMessages": ["to-d"], "dialogState": state("d", "say:D", "Exit")},
            {"id": 8, "internalSessionId": "sid-3", "time": 100, "message": "start", "botMessages": ["from-e"], "dialogState": state("e", "say:E", "Router")},
            {"id": 9, "internalSessionId": "sid-3", "time": 200, "message": "loop1", "botMessages": ["ask-a"], "dialogState": state("a", "say:A", "Main")},
            {"id": 10, "internalSessionId": "sid-3", "time": 300, "message": "loop2", "botMessages": ["retry-a"], "dialogState": state("a", "say:A", "Main")},
            {"id": 11, "internalSessionId": "sid-3", "time": 400, "message": "loop3", "botMessages": ["final-a"], "dialogState": state("a", "say:A", "Main")},
            {"id": 12, "internalSessionId": "sid-4", "time": 100, "message": "other", "botMessages": ["other-f"], "dialogState": state("f", "say:F", "Other")},
        ]

        with zipfile.ZipFile(zip_path, "w") as archive:
            archive.writestr(
                "messages.json",
                "".join(json.dumps(row) + "\n" for row in message_rows),
            )

    def test_build_block_drilldown_data_uses_block_id_as_primary_key(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            zip_path = Path(tmp_dir) / "conversations.zip"
            self._make_zip(zip_path)

            data = build_block_drilldown_data(zip_path, "a")

            self.assertEqual(data.target_block_id, "a")
            self.assertEqual(data.target_block_name, "say:A")
            self.assertEqual(data.target_context_name, "Main")
            self.assertEqual(data.total_sessions, 4)
            self.assertEqual(data.sessions_with_target_block, 3)
            self.assertEqual(data.total_target_rows, 6)
            self.assertEqual(data.total_target_runs, 3)
            self.assertEqual(data.runs_with_previous_block, 2)
            self.assertEqual(data.runs_with_next_block, 2)
            self.assertEqual(data.runs_with_fallback, 2)
            self.assertEqual(data.fallback_turns, 3)
            self.assertEqual(data.previous_user_messages, {"toA": 1, "start": 1})
            self.assertEqual(
                data.phrases_to_current,
                {
                    "hello": 1,
                    "still help": 1,
                    "question": 1,
                    "loop1": 1,
                    "loop2": 1,
                    "loop3": 1,
                },
            )
            self.assertEqual(data.phrases_to_next, {"still help": 1, "question": 1})
            self.assertEqual(data.fallback_phrases, {"hello": 1, "loop1": 1, "loop2": 1})
            self.assertEqual(data.previous_bot_messages, {"ask-a": 2, "welcome": 1, "from-c": 1, "from-e": 1, "retry-a": 1})
            self.assertEqual(data.next_bot_messages, {"ask-a": 3, "clarify-a": 1, "retry-a": 1, "final-a": 1})
            self.assertEqual(data.fallback_bot_messages, {"ask-a": 2, "retry-a": 1})
            self.assertEqual(data.next_blocks, {"b": 1, "d": 1})
            self.assertEqual(data.previous_blocks, {"c": 1, "e": 1})
            self.assertEqual(data.block_context_names["a"], "Main")

    def test_write_block_next_blocks_report_creates_one_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            base = Path(tmp_dir)
            zip_path = base / "conversations.zip"
            output_path = base / "next.xlsx"
            self._make_zip(zip_path)

            data = build_block_drilldown_data(zip_path, "a")
            write_block_next_blocks_report(data, output_path)

            workbook = load_workbook(output_path)
            self.assertEqual(workbook.sheetnames, ["Next Blocks"])
            rows = list(workbook.active.iter_rows(values_only=True))
            self.assertEqual(
                rows,
                [
                    ("Next block id", "Next block name", "Next context name", "Occurrences", "Percentage"),
                    ("b", "say:B", "Exit", 1, 0.5),
                    ("d", "say:D", "Exit", 1, 0.5),
                ],
            )

    def test_write_block_drilldown_report_creates_all_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            base = Path(tmp_dir)
            zip_path = base / "conversations.zip"
            output_path = base / "drilldown.xlsx"
            self._make_zip(zip_path)

            data = build_block_drilldown_data(zip_path, "a")
            write_block_drilldown_report(data, output_path)

            workbook = load_workbook(output_path)
            self.assertEqual(
                workbook.sheetnames,
                [
                    "Overview",
                    "Previous Blocks",
                    "Previous user messages",
                    "Bot messages",
                    "User replies",
                    "Next bot messages",
                    "Non-fallback user messages",
                    "Fallback user messages",
                    "Fallback bot replies",
                    "Next Blocks",
                ],
            )
            overview_rows = list(workbook["Overview"].iter_rows(values_only=True))
            self.assertIn(
                ("Target block id", "a", "Exact dialogState.id selected for this drilldown report."),
                overview_rows,
            )
            self.assertIn(
                ("Target context name", "Main", "Most frequent observed context name for the selected block id."),
                overview_rows,
            )
            self.assertIn(
                (
                    "Total target runs",
                    3,
                    "Consecutive groups of user turns with the selected dialogState.id after collapsing repeated adjacent turns.",
                ),
                overview_rows,
            )
            self.assertIn(
                (
                    "Runs with fallback",
                    2,
                    "Target runs whose length is greater than one, meaning the same input-requesting dialogState.id repeated across consecutive user turns. This is a fallback heuristic, not proof of exact runtime state.",
                ),
                overview_rows,
            )
            self.assertIn(
                (
                    "Unique previous user messages",
                    2,
                    "Different user messages that appeared just before this block became active.",
                ),
                overview_rows,
            )
            self.assertIn(
                (
                    "Unique bot messages",
                    5,
                    "Different bot messages that users replied to while they were in this block.",
                ),
                overview_rows,
            )
            self.assertIn(
                (
                    "Unique user replies",
                    6,
                    "Different user messages sent as replies to this block.",
                ),
                overview_rows,
            )
            self.assertIn(
                (
                    "Unique next bot messages",
                    4,
                    "Different bot messages sent after those user replies.",
                ),
                overview_rows,
            )
            self.assertIn(
                (
                    "Unique non-fallback user messages",
                    2,
                    "Different user replies that were followed by a different next block.",
                ),
                overview_rows,
            )
            self.assertIn(
                (
                    "Unique fallback user messages",
                    3,
                    "Different user replies after which the next exported user turn still referenced the same dialogState.id.",
                ),
                overview_rows,
            )
            self.assertIn(
                (
                    "Unique fallback bot replies",
                    2,
                    "Different bot replies emitted after user replies in repeated same-dialogState runs.",
                ),
                overview_rows,
            )
            self.assertEqual(
                list(workbook["Previous user messages"].iter_rows(values_only=True)),
                [
                    ("User message", "Occurrences"),
                    ("start", 1),
                    ("toA", 1),
                ],
            )
            self.assertEqual(
                list(workbook["Bot messages"].iter_rows(values_only=True)),
                [
                    ("Bot message", "Occurrences"),
                    ("ask-a", 2),
                    ("from-c", 1),
                    ("from-e", 1),
                    ("retry-a", 1),
                    ("welcome", 1),
                ],
            )
            self.assertEqual(
                list(workbook["User replies"].iter_rows(values_only=True)),
                [
                    ("User reply", "Occurrences"),
                    ("hello", 1),
                    ("loop1", 1),
                    ("loop2", 1),
                    ("loop3", 1),
                    ("question", 1),
                    ("still help", 1),
                ],
            )
            self.assertEqual(
                list(workbook["Next bot messages"].iter_rows(values_only=True)),
                [
                    ("Bot message", "Occurrences"),
                    ("ask-a", 3),
                    ("clarify-a", 1),
                    ("final-a", 1),
                    ("retry-a", 1),
                ],
            )

    def test_write_block_bot_messages_report_creates_one_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            base = Path(tmp_dir)
            zip_path = base / "conversations.zip"
            output_path = base / "bot-messages.xlsx"
            self._make_zip(zip_path)

            data = build_block_drilldown_data(zip_path, "a")
            write_block_bot_messages_report(data, output_path)

            workbook = load_workbook(output_path)
            self.assertEqual(workbook.sheetnames, ["Next bot messages"])
            self.assertEqual(
                list(workbook.active.iter_rows(values_only=True)),
                [
                    ("Bot message", "Occurrences"),
                    ("ask-a", 3),
                    ("clarify-a", 1),
                    ("final-a", 1),
                    ("retry-a", 1),
                ],
            )

    def test_write_block_previous_bot_messages_report_creates_one_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            base = Path(tmp_dir)
            zip_path = base / "conversations.zip"
            output_path = base / "previous-bot-messages.xlsx"
            self._make_zip(zip_path)

            data = build_block_drilldown_data(zip_path, "a")
            write_block_previous_bot_messages_report(data, output_path)

            workbook = load_workbook(output_path)
            self.assertEqual(workbook.sheetnames, ["Bot messages"])
            self.assertEqual(
                list(workbook.active.iter_rows(values_only=True)),
                [
                    ("Bot message", "Occurrences"),
                    ("ask-a", 2),
                    ("from-c", 1),
                    ("from-e", 1),
                    ("retry-a", 1),
                    ("welcome", 1),
                ],
            )

    def test_write_block_previous_user_messages_report_creates_one_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            base = Path(tmp_dir)
            zip_path = base / "conversations.zip"
            output_path = base / "previous-user-messages.xlsx"
            self._make_zip(zip_path)

            data = build_block_drilldown_data(zip_path, "a")
            write_block_previous_user_messages_report(data, output_path)

            workbook = load_workbook(output_path)
            self.assertEqual(workbook.sheetnames, ["Previous user messages"])
            self.assertEqual(
                list(workbook.active.iter_rows(values_only=True)),
                [
                    ("User message", "Occurrences"),
                    ("start", 1),
                    ("toA", 1),
                ],
            )


if __name__ == "__main__":
    unittest.main()
